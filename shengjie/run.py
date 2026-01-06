

from calibrator.Component.metrics import ECE, Accuracy, AdaptiveECE, ClasswiseECE, NLL, ECEDebiased, ECESweep, BrierLoss, RBS
import torch
import numpy as np
import json

# =========================
# [MOD] Helpers
# =========================
def _to_float(x):
    """Convert torch/numpy scalars to python float for JSON/Excel."""
    try:
        if isinstance(x, torch.Tensor):
            if x.numel() == 1:
                return float(x.detach().cpu().item())
            return float(x.detach().cpu().mean().item())
    except Exception:
        pass
    try:
        import numpy as _np
        if isinstance(x, (_np.generic,)):
            return float(x)
    except Exception:
        pass
    try:
        return float(x)
    except Exception:
        return x

def metric(logit, label, softmaxes=None, predictions=None, save=None):
    """
    Compute metrics and optionally save as JSON.
    [MOD] Now returns a python dict (floats) so caller can also write Excel easily.
    """
    if isinstance(logit, np.ndarray):
        logit = torch.from_numpy(logit)
    if isinstance(label, np.ndarray):
        label = torch.from_numpy(label)
    if isinstance(softmaxes, np.ndarray):
        softmaxes = torch.from_numpy(softmaxes)
    if predictions is not None and not torch.is_tensor(predictions):
        predictions = torch.as_tensor(predictions)

    device = "cuda"
    if logit is not None:
        logit = logit.to(device)
    if label is not None:
        label = label.to(device)
    if softmaxes is not None:
        softmaxes = softmaxes.to(device)
    if predictions is not None:
        predictions = predictions.to(device)

    # [MOD] compute once + convert to float
    ece = ECE().to(device)(logits=logit, labels=label, softmaxes=softmaxes)
    acc = Accuracy().to(device)(logits=logit, labels=label, softmaxes=softmaxes, predictions=predictions)
    adaece = AdaptiveECE().to(device)(logits=logit, labels=label, softmaxes=softmaxes)
    cece = ClasswiseECE().to(device)(logits=logit, labels=label, softmaxes=softmaxes)
    nll = NLL().to(device)(logits=logit, labels=label, softmaxes=softmaxes)
    ece_deb = ECEDebiased().to(device)(logits=logit, labels=label, softmaxes=softmaxes)
    ece_sweep = ECESweep().to(device)(logits=logit, labels=label, softmaxes=softmaxes)
    brier = BrierLoss().to(device)(logits=logit, labels=label, softmaxes=softmaxes)
    rbs = RBS().to(device)(logits=logit, labels=label, softmaxes=softmaxes)

    print("ECE: ", ece)
    print("Accuracy: ", acc)
    print("AdaptiveECE: ", adaece)
    print("ClasswiseECE: ", cece)
    print("NLL: ", nll)
    print("ECEDebiased: ", ece_deb)
    print("ECESweep: ", ece_sweep)
    print("BrierLoss: ", float(brier.detach().cpu()))
    print("RBS: ", float(rbs.detach().cpu()))

    result = {
        "ECE": _to_float(ece),
        "Accuracy": _to_float(acc),
        "AdaptiveECE": _to_float(adaece),
        "ClasswiseECE": _to_float(cece),
        "NLL": _to_float(nll),
        "ECEDebiased": _to_float(ece_deb),
        "ECESweep": _to_float(ece_sweep),
        "BrierLoss": float(brier.detach().cpu()),
        "RBS": float(rbs.detach().cpu()),
    }

    if save is not None:
        with open(save, "w", encoding="utf-8") as f:
            json.dump(result, f, ensure_ascii=False, indent=2)

    return result




from calibrator.Component import TemperatureScalingCalibrator, PTSCalibrator, CTSCalibrator, ETSCalibrator, HistogramBinningCalibrator
from calibrator.Component import BBQCalibrator, VectorScalingCalibrator, GroupCalibrationCalibrator, ProCalDensityRatioCalibrator
from calibrator.Component import SplineCalibrator
from calibrator.Component.model.feature_clipping import FeatureClippingCalibrator
from utils.smart_calibrator import SMART
import os
import re  # [MOD]
from types import SimpleNamespace
from utils.model_utils import create_model
from sklearn.utils import resample
from typing import Dict, Optional  # [MOD]


# =========================
# [MOD] Output path helpers
# =========================
def _project_root():
    return os.path.dirname(os.path.abspath(__file__))

def _result_root():
    return os.path.join(_project_root(), "result")

def _ensure_dir(d: str):
    if d:
        os.makedirs(d, exist_ok=True)

def _build_case_dir(dataset: str, model: str, seed: int, vs: float, bs: Optional[float], bootstrap: bool) -> str:
    """
    result/<dataset>_<model>/<case_name>/
    case_name:
      - seed1_vs0.2
      - seed1_vs0.2_bs0.1
    """
    dataset_model_dir = os.path.join(_result_root(), f"{dataset}_{model}")
    _ensure_dir(dataset_model_dir)

    if bootstrap:
        case_name = f"seed{seed}_vs{vs}_bs{bs}"
    else:
        case_name = f"seed{seed}_vs{vs}"
    case_dir = os.path.join(dataset_model_dir, case_name)
    _ensure_dir(case_dir)
    return case_dir

def _excel_path(case_dir: str, dataset: str, model: str, seed: int, vs: float, bs: Optional[float], bootstrap: bool) -> str:
    if bootstrap:
        name = f"{dataset}_{model}_seed{seed}_vs{vs}_bs{bs}.xlsx"
    else:
        name = f"{dataset}_{model}_seed{seed}_vs{vs}.xlsx"
    return os.path.join(case_dir, name)

def _pt_for_seed(pt_template: str, seed: int) -> str:
    """
    [MOD] Derive pt path for seed=1..5 from a template path.
    Supported:
      1) '...seed1...' -> replace number
      2) '...{seed}...' -> format
    """
    if "{seed}" in pt_template:
        try:
            return pt_template.format(seed=seed)
        except Exception:
            pass
    m = re.search(r"seed(\d+)", pt_template)
    if not m:
        return pt_template
    start, end = m.span(1)
    return pt_template[:start] + str(seed) + pt_template[end:]


# =========================
# Methods (core logic unchanged; only output routing + return dict)
# =========================
def Uncalibrated(pt, output_dir: Optional[str] = None):
    val_logits, val_features, val_labels, test_logits, test_features, test_labels = torch.load(pt, weights_only=False)
    logit = test_logits
    label = test_labels
    print(pt)
    print('Uncalibrated')

    # [MOD] output routing
    if output_dir is None:
        OUTPUT_DIR = _result_root()
        save_path = os.path.join(OUTPUT_DIR, f"{pt}_01Uncalibrated.json")
        _ensure_dir(os.path.dirname(save_path))
    else:
        _ensure_dir(output_dir)
        save_path = os.path.join(output_dir, "Uncalibrated.json")

    return metric(logit, label, save=save_path)

def TS(pt, seed=None, bs=None, vs=None, bootstrap=False, output_dir: Optional[str] = None):
    val_logits, val_features, val_labels, test_logits, test_features, test_labels = torch.load(pt, weights_only=False)
    if bootstrap:
        val_logits, val_features, val_labels = resample(
            val_logits, val_features, val_labels,
            replace=True,
            random_state=seed,
            n_samples=int(round((bs / vs) * len(val_labels)))
        )
    calibrator = TemperatureScalingCalibrator()
    calibrator.fit(val_logits, val_labels)
    calibrated_logits = calibrator.calibrate(test_logits, return_logits=True)
    logit = calibrated_logits
    label = test_labels
    print(pt)
    print('TS')

    if output_dir is None:
        OUTPUT_DIR = _result_root()
        save_path = os.path.join(OUTPUT_DIR, f"{pt}_02TS_{bs}.json" if bootstrap else f"{pt}_02TS.json")
        _ensure_dir(os.path.dirname(save_path))
    else:
        _ensure_dir(output_dir)
        save_path = os.path.join(output_dir, "TS.json")

    return metric(logit, label, save=save_path)

def PTS(pt, seed, bs=None, vs=None, bootstrap=False, output_dir: Optional[str] = None):
    val_logits, val_features, val_labels, test_logits, test_features, test_labels = torch.load(pt, weights_only=False)
    if bootstrap:
        val_logits, val_features, val_labels = resample(
            val_logits, val_features, val_labels,
            replace=True,
            random_state=seed,
            n_samples=int(round((bs / vs) * len(val_labels)))
        )
    num_classes = val_logits.shape[1]
    calibrator = PTSCalibrator(
        length_logits=num_classes,
        steps=10000,
        lr=0.00005,
        nlayers=2,
        n_nodes=5,
        loss_fn="mse",
        top_k_logits=10,
        seed=seed
    )
    calibrator.fit(val_logits, val_labels)
    calibrated_logits = calibrator.calibrate(test_logits, return_logits=True)
    logit = calibrated_logits
    label = test_labels
    print(pt)
    print('PTS')

    if output_dir is None:
        OUTPUT_DIR = _result_root()
        save_path = os.path.join(OUTPUT_DIR, f"{pt}_03PTS_{bs}.json" if bootstrap else f"{pt}_03PTS.json")
        _ensure_dir(os.path.dirname(save_path))
    else:
        _ensure_dir(output_dir)
        save_path = os.path.join(output_dir, "PTS.json")

    return metric(logit, label, save=save_path)

def CTS(pt, seed=None, bs=None, vs=None, bootstrap=False, output_dir: Optional[str] = None):
    val_logits, val_features, val_labels, test_logits, test_features, test_labels = torch.load(pt, weights_only=False)
    if bootstrap:
        val_logits, val_features, val_labels = resample(
            val_logits, val_features, val_labels,
            replace=True,
            random_state=seed,
            n_samples=int(round((bs / vs) * len(val_labels)))
        )
    num_classes = val_logits.shape[1]
    calibrator = CTSCalibrator(
        n_class=num_classes,
        n_bins=15,
        n_iter=5
    )
    calibrator.fit(val_logits, val_labels)
    calibrated_logits = calibrator.calibrate(test_logits, return_logits=True)
    logit = calibrated_logits
    label = test_labels
    print(pt)
    print('CTS')

    if output_dir is None:
        OUTPUT_DIR = _result_root()
        save_path = os.path.join(OUTPUT_DIR, f"{pt}_04CTS_{bs}.json" if bootstrap else f"{pt}_04CTS.json")
        _ensure_dir(os.path.dirname(save_path))
    else:
        _ensure_dir(output_dir)
        save_path = os.path.join(output_dir, "CTS.json")

    return metric(logit, label, save=save_path)

def ETS(pt, seed=None, bs=None, vs=None, bootstrap=False, output_dir: Optional[str] = None):
    val_logits, val_features, val_labels, test_logits, test_features, test_labels = torch.load(pt, weights_only=False)
    if bootstrap:
        val_logits, val_features, val_labels = resample(
            val_logits, val_features, val_labels,
            replace=True,
            random_state=seed,
            n_samples=int(round((bs / vs) * len(val_labels)))
        )
    num_classes = val_logits.shape[1]
    calibrator = ETSCalibrator(n_classes=num_classes)
    calibrator.fit(val_logits, val_labels)
    calibrated_softmaxes = calibrator.calibrate(test_logits, return_logits=False)
    softmaxes = calibrated_softmaxes
    label = test_labels
    print(pt)
    print('ETS')

    if output_dir is None:
        OUTPUT_DIR = _result_root()
        save_path = os.path.join(OUTPUT_DIR, f"{pt}_05ETS_{bs}.json" if bootstrap else f"{pt}_05ETS.json")
        _ensure_dir(os.path.dirname(save_path))
    else:
        _ensure_dir(output_dir)
        save_path = os.path.join(output_dir, "ETS.json")

    return metric(logit=None, label=label, softmaxes=softmaxes, save=save_path)

def SMART_(pt, dataset, model, seed, vs, loss, bs=None, bootstrap=False, output_dir: Optional[str] = None):
    val_logits, val_features, val_labels, test_logits, test_features, test_labels = torch.load(pt, weights_only=False)
    if bootstrap:
        val_logits, val_features, val_labels = resample(
            val_logits, val_features, val_labels,
            replace=True,
            random_state=seed,
            n_samples=int(round((bs / vs) * len(val_labels)))
        )

    args = SimpleNamespace(dataset_root='data')
    smart_epochs = 200
    dataset_name = dataset
    model_name = model
    seed_value = seed
    # valid_size = vs
    valid_size = bs
    smart_loss = getattr(args, 'smart_loss', 'soft_ece')
    patience = 20
    min_delta = 0.0001
    corruption_type = getattr(args, 'corruption_type', None) if dataset_name == 'imagenet_c' else None
    severity = getattr(args, 'severity', None) if dataset_name == 'imagenet_c' else None
    train_loss = loss

    smart = SMART(
        epochs=smart_epochs, dataset_name=dataset_name,
        model_name=model_name, seed_value=seed_value,
        valid_size=valid_size, loss_fn=smart_loss,
        patience=patience, min_delta=min_delta,
        corruption_type=corruption_type, severity=severity,
        train_loss=train_loss
    )
    smart.fit(val_logits, val_labels)
    smart_logits = smart.calibrate(test_logits, return_logits=True)
    logit = smart_logits
    label = test_labels
    print(pt)
    print('SMART')

    if output_dir is None:
        OUTPUT_DIR = _result_root()
        save_path = os.path.join(OUTPUT_DIR, f"{pt}_06SMART_{bs}.json" if bootstrap else f"{pt}_06SMART.json")
        _ensure_dir(os.path.dirname(save_path))
    else:
        _ensure_dir(output_dir)
        save_path = os.path.join(output_dir, "SMART.json")

    return metric(logit, label, save=save_path)

def HB(pt, seed=None, bs=None, vs=None, bootstrap=False, output_dir: Optional[str] = None):
    val_logits, val_features, val_labels, test_logits, test_features, test_labels = torch.load(pt, weights_only=False)
    if bootstrap:
        val_logits, val_features, val_labels = resample(
            val_logits, val_features, val_labels,
            replace=True,
            random_state=seed,
            n_samples=int(round((bs / vs) * len(val_labels)))
        )
    calibrator = HistogramBinningCalibrator(n_bins=15, strategy='uniform')
    calibrator.fit(val_logits, val_labels)
    calibrated_logits = calibrator.calibrate(test_logits, return_logits=True)
    logit = calibrated_logits
    label = test_labels
    print(pt)
    print('HB')

    if output_dir is None:
        OUTPUT_DIR = _result_root()
        save_path = os.path.join(OUTPUT_DIR, f"{pt}_07HB_{bs}.json" if bootstrap else f"{pt}_07HB.json")
        _ensure_dir(os.path.dirname(save_path))
    else:
        _ensure_dir(output_dir)
        save_path = os.path.join(output_dir, "HB.json")

    return metric(logit=logit, label=label, save=save_path)

def BBQ(pt, seed=None, bs=None, vs=None, bootstrap=False, output_dir: Optional[str] = None):
    val_logits, val_features, val_labels, test_logits, test_features, test_labels = torch.load(pt, weights_only=False)
    if bootstrap:
        val_logits, val_features, val_labels = resample(
            val_logits, val_features, val_labels,
            replace=True,
            random_state=seed,
            n_samples=int(round((bs / vs) * len(val_labels)))
        )
    calibrator = BBQCalibrator(score_type='max_prob', n_bins_max=20)
    calibrator.fit(val_logits, val_labels)
    calibrated_softmaxes = calibrator.calibrate(test_logits, return_logits=False)
    softmaxes = calibrated_softmaxes
    label = test_labels
    print(pt)
    print('BBQ')

    if output_dir is None:
        OUTPUT_DIR = _result_root()
        save_path = os.path.join(OUTPUT_DIR, f"{pt}_08BBQ_{bs}.json" if bootstrap else f"{pt}_08BBQ.json")
        _ensure_dir(os.path.dirname(save_path))
    else:
        _ensure_dir(output_dir)
        save_path = os.path.join(output_dir, "BBQ.json")

    return metric(logit=None, label=label, softmaxes=softmaxes, save=save_path)

def VS(pt, seed=None, bs=None, vs=None, bootstrap=False, output_dir: Optional[str] = None):
    val_logits, val_features, val_labels, test_logits, test_features, test_labels = torch.load(pt, weights_only=False)
    if bootstrap:
        val_logits, val_features, val_labels = resample(
            val_logits, val_features, val_labels,
            replace=True,
            random_state=seed,
            n_samples=int(round((bs / vs) * len(val_labels)))
        )
    calibrator = VectorScalingCalibrator(loss_type='nll', bias=True)
    calibrator.fit(val_logits, val_labels)
    calibrated_logits = calibrator.calibrate(test_logits, return_logits=True)
    logit = calibrated_logits
    label = test_labels
    print(pt)
    print('VS')

    if output_dir is None:
        OUTPUT_DIR = _result_root()
        save_path = os.path.join(OUTPUT_DIR, f"{pt}_09VS_{bs}.json" if bootstrap else f"{pt}_09VS.json")
        _ensure_dir(os.path.dirname(save_path))
    else:
        _ensure_dir(output_dir)
        save_path = os.path.join(output_dir, "VS.json")

    return metric(logit, label, save=save_path)

def GC(pt, seed=None, bs=None, vs=None, bootstrap=False, output_dir: Optional[str] = None):
    val_logits, val_features, val_labels, test_logits, test_features, test_labels = torch.load(pt, weights_only=False)
    if bootstrap:
        val_logits, val_features, val_labels = resample(
            val_logits, val_features, val_labels,
            replace=True,
            random_state=seed,
            n_samples=int(round((bs / vs) * len(val_labels)))
        )
    calibrator = GroupCalibrationCalibrator(
        num_groups=2,
        num_partitions=20,
        weight_decay=0.1
    )
    calibrator.fit(val_logits, val_labels)
    calibrated_logits = calibrator.calibrate(test_logits, return_logits=True)
    logit = calibrated_logits
    label = test_labels
    print(pt)
    print('GC')

    if output_dir is None:
        OUTPUT_DIR = _result_root()
        save_path = os.path.join(OUTPUT_DIR, f"{pt}_10GC_{bs}.json" if bootstrap else f"{pt}_10GC.json")
        _ensure_dir(os.path.dirname(save_path))
    else:
        _ensure_dir(output_dir)
        save_path = os.path.join(output_dir, "GC.json")

    return metric(logit, label, save=save_path)

def ProCal_DR(pt, seed=None, bs=None, vs=None, bootstrap=False, output_dir: Optional[str] = None):
    val_logits, val_features, val_labels, test_logits, test_features, test_labels = torch.load(pt, weights_only=False)
    if bootstrap:
        val_logits, val_features, val_labels = resample(
            val_logits, val_features, val_labels,
            replace=True,
            random_state=seed,
            n_samples=int(round((bs / vs) * len(val_labels)))
        )
    calibrator = ProCalDensityRatioCalibrator(
        k_neighbors=10,
        bandwidth='normal_reference',
        kernel='KDEMultivariate',
        distance_measure='L2',
        normalize_features=True
    )
    calibrator.fit(val_logits, val_labels, val_features)
    calibrated_softmaxes = calibrator.calibrate(test_logits, test_features, return_logits=False)
    softmaxes = calibrated_softmaxes
    label = test_labels
    print(pt)
    print('ProCal_DR')

    if output_dir is None:
        OUTPUT_DIR = _result_root()
        save_path = os.path.join(OUTPUT_DIR, f"{pt}_11ProCal_DR_{bs}.json" if bootstrap else f"{pt}_11ProCal_DR.json")
        _ensure_dir(os.path.dirname(save_path))
    else:
        _ensure_dir(output_dir)
        save_path = os.path.join(output_dir, "ProCal_DR.json")

    return metric(logit=None, label=label, softmaxes=softmaxes, save=save_path)

def FC(pt, model, dataset, seed=None, bs=None, vs=None, bootstrap=False, output_dir: Optional[str] = None):
    val_logits, val_features, val_labels, test_logits, test_features, test_labels = torch.load(pt, weights_only=False)
    if bootstrap:
        val_logits, val_features, val_labels = resample(
            val_logits, val_features, val_labels,
            replace=True,
            random_state=seed,
            n_samples=int(round((bs / vs) * len(val_labels)))
        )
    args = SimpleNamespace(dataset_root='data')
    model_name = model
    dataset_name = dataset
    device = 'cuda'
    val_logits = val_logits.to(device)
    val_features = val_features.to(device)
    val_labels = val_labels.to(device)
    test_logits = test_logits.to(device)
    test_features = test_features.to(device)
    test_labels = test_labels.to(device)

    model_obj = create_model(args, model_name, dataset_name, device)
    model_obj.eval()
    classifier_fn = model_obj.classifier

    calibrator = FeatureClippingCalibrator(cross_validate='ece')
    optimal_clip = calibrator.set_feature_clip(val_features, val_logits, val_labels, classifier_fn)
    clipped_test_features = calibrator.feature_clipping(test_features, optimal_clip)
    fc_logits = classifier_fn(clipped_test_features)
    logit = fc_logits
    label = test_labels
    print(pt)
    print('FC')

    if output_dir is None:
        OUTPUT_DIR = _result_root()
        save_path = os.path.join(OUTPUT_DIR, f"{pt}_12FC_{bs}.json" if bootstrap else f"{pt}_12FC.json")
        _ensure_dir(os.path.dirname(save_path))
    else:
        _ensure_dir(output_dir)
        save_path = os.path.join(output_dir, "FC.json")

    return metric(logit, label, save=save_path)

def Spline(pt, seed=None, bs=None, vs=None, bootstrap=False, output_dir: Optional[str] = None):
    val_logits, val_features, val_labels, test_logits, test_features, test_labels = torch.load(pt, weights_only=False)
    if bootstrap:
        val_logits, val_features, val_labels = resample(
            val_logits, val_features, val_labels,
            replace=True,
            random_state=seed,
            n_samples=int(round((bs / vs) * len(val_labels)))
        )
    calibrator = SplineCalibrator()
    calibrator.fit(val_logits, val_labels)
    calibrated_softmaxes = calibrator.calibrate(test_logits)
    softmaxes = calibrated_softmaxes
    label = test_labels
    print(pt)
    print('Spline')

    if output_dir is None:
        OUTPUT_DIR = _result_root()
        save_path = os.path.join(OUTPUT_DIR, f"{pt}_13Spline_{bs}.json" if bootstrap else f"{pt}_13Spline.json")
        _ensure_dir(os.path.dirname(save_path))
    else:
        _ensure_dir(output_dir)
        save_path = os.path.join(output_dir, "Spline.json")

    return metric(logit=None, label=label, softmaxes=softmaxes, save=save_path)


# =========================
# [MOD] Excel writer
# =========================
EXCEL_COLUMNS = [
    "PostHoc_Method", "ECE", "Accuracy", "AdaECE", "CECE", "NLL",
    "ECE_debiased", "ECE_sweep", "Brier", "RBS", "Notes"
]

METHOD_ORDER = [
    "Uncalibrated", "TS", "PTS", "CTS", "ETS", "SMART",
    "HB", "BBQ", "VS", "GC", "ProCal_DR", "FC", "Spline"
]

def _write_excel(excel_file: str, results: Dict[str, Optional[dict]], notes: Dict[str, str]):
    """
    Create an Excel file with the template shown in the screenshot.
    - If a method didn't run: keep row but empty.
    - If a method crashed: keep row, leave metrics empty, put error in Notes.
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment, Border, Side
    except Exception as e:
        raise RuntimeError("openpyxl is required to write Excel. Please `pip install openpyxl`.") from e

    wb = Workbook()
    ws = wb.active
    ws.title = "results"

    # header
    ws.append(EXCEL_COLUMNS)
    header_font = Font(bold=True)
    for col in range(1, len(EXCEL_COLUMNS) + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # rows
    for m in METHOD_ORDER:
        r = results.get(m)
        n = notes.get(m, "")

        if r is None:
            ws.append([m, "", "", "", "", "", "", "", "", "", n])
            continue

        ws.append([
            m,
            r.get("ECE", ""),
            r.get("Accuracy", ""),
            r.get("AdaptiveECE", ""),   # AdaECE
            r.get("ClasswiseECE", ""),  # CECE
            r.get("NLL", ""),
            r.get("ECEDebiased", ""),
            r.get("ECESweep", ""),
            r.get("BrierLoss", ""),
            r.get("RBS", ""),
            n
        ])

    # basic grid/borders like the screenshot
    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=len(EXCEL_COLUMNS)):
        for cell in row:
            cell.border = border
            if cell.row > 1:
                cell.alignment = Alignment(vertical="center")

    # column widths
    widths = {
        "A": 18, "B": 12, "C": 12, "D": 12, "E": 12, "F": 12,
        "G": 14, "H": 12, "I": 12, "J": 12, "K": 60
    }
    for k, v in widths.items():
        ws.column_dimensions[k].width = v

    _ensure_dir(os.path.dirname(excel_file))
    wb.save(excel_file)


# =========================
# [MOD] Batch runner
# =========================
BS_LIST = [0.1, 0.05, 0.01, 0.001, 0.0002]
SEEDS = [1, 2, 3, 4, 5]

def _save_error_json(case_dir: str, method_name: str, err: str):
    try:
        p = os.path.join(case_dir, f"{method_name}.json")
        with open(p, "w", encoding="utf-8") as f:
            json.dump({"error": err}, f, ensure_ascii=False, indent=2)
    except Exception:
        pass

def evaluate(pt=None, seed=None, dataset=None, model=None, vs=None, loss=None, bs=None, bootstrap=None):
    """
    [MOD] New behavior:
    - For the input pt (assumed to contain 'seed1' etc), automatically run:
        seeds 1..5
          - bootstrap=False
          - bootstrap=True with bs in [0.1, 0.05, 0.01, 0.001, 0.0002]
      => 30 cases per pt base.
    - Directory layout:
        result/<dataset>_<model>/<case_folder>/
          - <Method>.json
          - <dataset>_<model>_seedX_vsY[_bsZ].xlsx
    - Error handling:
        if a method crashes -> skip it, keep row, write error to Notes.
    """
    # keep signature unchanged, but `seed/bs/bootstrap` are ignored because we run the full grid.
    if pt is None or dataset is None or model is None or vs is None:
        raise ValueError("evaluate requires pt, dataset, model, vs (loss is needed for SMART).")

    # choose which methods to run (edit this list if you want to skip some)
    methods_to_run = set(METHOD_ORDER)

    for s in SEEDS:
        pt_s = _pt_for_seed(pt, s)

        # case 1) bootstrap=False
        for bootstrap_flag, bs_val in [(False, None)] + [(True, b) for b in BS_LIST]:
            case_dir = _build_case_dir(dataset, model, s, vs, bs_val, bootstrap_flag)

            results: Dict[str, Optional[dict]] = {m: None for m in METHOD_ORDER}
            notes: Dict[str, str] = {m: "" for m in METHOD_ORDER}

            # runners for this case
            def _run(method_name: str):
                if method_name == "Uncalibrated":
                    return Uncalibrated(pt_s, output_dir=case_dir)
                if method_name == "TS":
                    return TS(pt_s, seed=s, bs=bs_val, vs=vs, bootstrap=bootstrap_flag, output_dir=case_dir)
                if method_name == "PTS":
                    return PTS(pt_s, seed=s, bs=bs_val, vs=vs, bootstrap=bootstrap_flag, output_dir=case_dir)
                if method_name == "CTS":
                    return CTS(pt_s, seed=s, bs=bs_val, vs=vs, bootstrap=bootstrap_flag, output_dir=case_dir)
                if method_name == "ETS":
                    return ETS(pt_s, seed=s, bs=bs_val, vs=vs, bootstrap=bootstrap_flag, output_dir=case_dir)
                if method_name == "SMART":
                    # [MOD] SMART_ uses valid_size=bs internally; for non-bootstrap we pass bs=vs to avoid None.
                    smart_bs = bs_val if bootstrap_flag else vs
                    return SMART_(pt_s, dataset, model, s, vs, loss, bs=smart_bs, bootstrap=bootstrap_flag, output_dir=case_dir)
                if method_name == "HB":
                    return HB(pt_s, seed=s, bs=bs_val, vs=vs, bootstrap=bootstrap_flag, output_dir=case_dir)
                if method_name == "BBQ":
                    return BBQ(pt_s, seed=s, bs=bs_val, vs=vs, bootstrap=bootstrap_flag, output_dir=case_dir)
                if method_name == "VS":
                    return VS(pt_s, seed=s, bs=bs_val, vs=vs, bootstrap=bootstrap_flag, output_dir=case_dir)
                if method_name == "GC":
                    return GC(pt_s, seed=s, bs=bs_val, vs=vs, bootstrap=bootstrap_flag, output_dir=case_dir)
                if method_name == "ProCal_DR":
                    return ProCal_DR(pt_s, seed=s, bs=bs_val, vs=vs, bootstrap=bootstrap_flag, output_dir=case_dir)
                if method_name == "FC":
                    return FC(pt_s, model, dataset, seed=s, bs=bs_val, vs=vs, bootstrap=bootstrap_flag, output_dir=case_dir)
                # if method_name == "Spline":
                #     return Spline(pt_s, seed=s, bs=bs_val, vs=vs, bootstrap=bootstrap_flag, output_dir=case_dir)
                raise ValueError(f"Unknown method: {method_name}")

            # run methods one by one (method-level try/except)
            for m in METHOD_ORDER:
                if m not in methods_to_run:
                    continue
                try:
                    results[m] = _run(m)
                except Exception as e:
                    err = f"{type(e).__name__}: {e}"
                    notes[m] = err
                    results[m] = None
                    _save_error_json(case_dir, m, err)
                finally:
                    try:
                        torch.cuda.empty_cache()
                    except Exception:
                        pass

            # write excel for this case
            xlsx = _excel_path(case_dir, dataset, model, s, vs, bs_val, bootstrap_flag)
            _write_excel(xlsx, results, notes)


if __name__ == '__main__':
    # evaluate(
    #     pt="output/imagenet_sketch_resnet50_seed1_vs0.2.pt",
    #     seed=1, dataset="imagenet_sketch", model="resnet50",
    #     vs=0.2, loss="ce", bs=0.1, bootstrap=True
    # )
    # evaluate(
    #     pt="output/imagenet_sketch_densenet121_seed1_vs0.2.pt",
    #     seed=1, dataset="imagenet_sketch", model="densenet121",
    #     vs=0.2, loss="ce", bs=0.1, bootstrap=True
    # )
    # evaluate(
    #     pt="output/imagenet_sketch_resnet152_seed1_vs0.2.pt",
    #     seed=1, dataset="imagenet_sketch", model="resnet152",
    #     vs=0.2, loss="ce", bs=0.1, bootstrap=True
    # )
    # evaluate(
    #     pt="output/imagenet_sketch_swin_b_seed1_vs0.2.pt",
    #     seed=1, dataset="imagenet_sketch", model="swin_b",
    #     vs=0.2, loss="ce", bs=0.1, bootstrap=True
    # )
    # evaluate(
    #     pt="output/imagenet_sketch_vit_b16_seed1_vs0.2.pt",
    #     seed=1, dataset="imagenet_sketch", model="vit_b_16",
    #     vs=0.2, loss="ce", bs=0.1, bootstrap=True
    # )
    # evaluate(
    #     pt="output/imagenet_sketch_vit_b32_seed1_vs0.2.pt",
    #     seed=1, dataset="imagenet_sketch", model="vit_b_32",
    #     vs=0.2, loss="ce", bs=0.1, bootstrap=True
    # )
    # evaluate(
    #     pt="output/imagenet_sketch_vit_l16_seed1_vs0.2.pt",
    #     seed=1, dataset="imagenet_sketch", model="vit_l_16",
    #     vs=0.2, loss="ce", bs=0.1, bootstrap=True
    # )
    # evaluate(
    #     pt="output/imagenet_sketch_wide_resnet50_seed1_vs0.2.pt",
    #     seed=1, dataset="imagenet_sketch", model="wide_resnet",
    #     vs=0.2, loss="ce", bs=0.1, bootstrap=True
    # )
    # evaluate(
    #     pt="output/imagenet_sketch_mobilenet_v2_seed1_vs0.2.pt",
    #     seed=1, dataset="imagenet_sketch", model="mobilenet_v2",
    #     vs=0.2, loss="ce", bs=0.1, bootstrap=True
    # )
    # evaluate(
    #     pt="output/imagenet_sketch_convnext_tiny_seed1_vs0.2.pt",
    #     seed=1, dataset="imagenet_sketch", model="convnext_tiny",
    #     vs=0.2, loss="ce", bs=0.1, bootstrap=True
    # )
    # evaluate(
    #     pt="output/imagenet_sketch_convnext_base_seed1_vs0.2.pt",
    #     seed=1, dataset="imagenet_sketch", model="convnext_base",
    #     vs=0.2, loss="ce", bs=0.1, bootstrap=True
    # )
    # evaluate(
    #     pt="output/imagenet_sketch_convnext_large_seed1_vs0.2.pt",
    #     seed=1, dataset="imagenet_sketch", model="convnext_large",
    #     vs=0.2, loss="ce", bs=0.1, bootstrap=True
    # )
    # evaluate(
    #     pt="output/imagenet_sketch_mlp_mixer_b16_seed1_vs0.2.pt",
    #     seed=1, dataset="imagenet_sketch", model="mlp_mixer_b16",
    #     vs=0.2, loss="ce", bs=0.1, bootstrap=True
    # )
    evaluate(
        pt="output/imagenet_sketch_beit_base_seed1_vs0.2.pt",
        seed=1, dataset="imagenet_sketch", model="beit_base",
        vs=0.2, loss="ce", bs=0.1, bootstrap=True
    )
    evaluate(
        pt="output/imagenet_sketch_beit_large_seed1_vs0.2.pt",
        seed=1, dataset="imagenet_sketch", model="beit_large",
        vs=0.2, loss="ce", bs=0.1, bootstrap=True
    )
    evaluate(
        pt="output/imagenet_sketch_beitv2_base_seed1_vs0.2.pt",
        seed=1, dataset="imagenet_sketch", model="beitv2_base",
        vs=0.2, loss="ce", bs=0.1, bootstrap=True
    )
    evaluate(
        pt="output/imagenet_sketch_eva02_small_seed1_vs0.2.pt",
        seed=1, dataset="imagenet_sketch", model="eva02_small",
        vs=0.2, loss="ce", bs=0.1, bootstrap=True
    )
    evaluate(
        pt="output/imagenet_sketch_eva02_base_seed1_vs0.2.pt",
        seed=1, dataset="imagenet_sketch", model="eva02_base",
        vs=0.2, loss="ce", bs=0.1, bootstrap=True
    )
    evaluate(
        pt="output/imagenet_sketch_eva02_large_seed1_vs0.2.pt",
        seed=1, dataset="imagenet_sketch", model="eva02_large",
        vs=0.2, loss="ce", bs=0.1, bootstrap=True
    )